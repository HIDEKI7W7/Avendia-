import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Colores de la plataforma (Brand Colors)
AZUL_INSTITUCIONAL = RGBColor(74, 144, 226)      # #4A90E2
MORADO_IA = RGBColor(124, 108, 242)             # #7C6CF2
GRIS_BORDE = "D3D3D3"
BG_HEADER_HEX = "4A90E2"                          # Azul Institucional para cabeceras
BG_DATOS_HEX = "F0F4FF"                           # Un azul extra claro para destacar datos informativos

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos (padding) en dxa para una celda (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Colorea el fondo de una celda con un color hexadecimal."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table, color=GRIS_BORDE, sz="4", val="single"):
    """Aplica bordes delgados de un color gris claro a toda la tabla."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_col_widths(table, widths):
    """Establece anchos exactos para cada columna de una tabla."""
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = width

def add_main_title(doc, text):
    """Agrega el título principal del documento centrado y estilizado."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = AZUL_INSTITUCIONAL
    return p

def add_section_heading(doc, text):
    """Agrega un encabezado de sección con color institucional y espaciado óptimo."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = AZUL_INSTITUCIONAL
    return p

def add_key_val_p(doc, key, val, bullet=False, space_after=4):
    """Agrega un párrafo formateado con la clave en negrita y el valor en regular."""
    style_name = 'List Bullet' if bullet else 'Normal'
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    
    run_key = p.add_run(key)
    run_key.bold = True
    run_key.font.name = 'Calibri'
    run_key.font.size = Pt(11)
    
    run_val = p.add_run(val)
    run_val.font.name = 'Calibri'
    run_val.font.size = Pt(11)
    return p

def fill_cell(cell, text, bold=False, italic=False, bg_hex=None, text_color_rgb=None, size_pt=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Limpia la celda e inyecta texto formateado párrafo por párrafo."""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.text = ""
    
    lines = text.split('\n')
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(0)
            
        run = p.add_run(line)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Calibri'
        run.font.size = Pt(size_pt)
        if text_color_rgb:
            run.font.color.rgb = text_color_rgb
            
    if bg_hex:
        set_cell_background(cell, bg_hex)
        
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

def generate_clase_template(output_path):
    """Genera la Plantilla 1: Planificación de Clase."""
    doc = Document()
    
    # Configurar márgenes de página (1 pulgada = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Título Principal
    add_main_title(doc, "PLANIFICACIÓN DE LA SESIÓN DE APRENDIZAJE")

    # I. DATOS INFORMATIVOS
    add_section_heading(doc, "I. DATOS INFORMATIVOS")
    
    # Tabla elegante para Datos Informativos (4 filas x 2 columnas)
    info_table = doc.add_table(rows=4, cols=2)
    set_table_borders(info_table)
    set_col_widths(info_table, [Inches(3.25), Inches(3.25)])
    
    # Llenar datos informativos
    fill_cell(info_table.cell(0, 0), "Institución Educativa: [Insertar Nombre de la I.E.]", bold=True)
    fill_cell(info_table.cell(0, 1), "Director(a): [Nombre del Director]", bold=True)
    fill_cell(info_table.cell(1, 0), "Docente: [Nombre del Docente]", bold=True)
    fill_cell(info_table.cell(1, 1), "Grado y Sección: [Ej. 3° Grado \"A\"]", bold=True)
    fill_cell(info_table.cell(2, 0), "Área Curricular: [Ej. Comunicación / Matemática]", bold=True)
    fill_cell(info_table.cell(2, 1), "Fecha: [Día / Mes / Año]", bold=True)
    fill_cell(info_table.cell(3, 0), "Duración: [Ej. 90 minutos]", bold=True)
    fill_cell(info_table.cell(3, 1), "", bold=True) # Celda vacía

    # II. TÍTULO DE LA SESIÓN
    add_section_heading(doc, "II. TÍTULO DE LA SESIÓN")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('[Escribir un título breve y motivador que sintetice la actividad principal, ej: "Leemos un texto interactivo sobre el cuidado del agua"]')
    run.font.italic = True
    run.font.name = 'Calibri'
    run.font.size = Pt(11)

    # III. PROPÓSITOS Y EVIDENCIAS DE APRENDIZAJE
    add_section_heading(doc, "III. PROPÓSITOS Y EVIDENCIAS DE APRENDIZAJE")
    
    # Tabla de 5 columnas para la matriz de aprendizaje
    # Total ancho = 6.5 pulgadas (1.4, 1.3, 1.3, 1.3, 1.2)
    table_prop = doc.add_table(rows=2, cols=5)
    set_table_borders(table_prop)
    set_col_widths(table_prop, [Inches(1.4), Inches(1.3), Inches(1.3), Inches(1.3), Inches(1.2)])
    
    # Cabecera
    headers = [
        "Competencia(s) /\nCapacidad(es)",
        "Desempeños\n(CNEB)",
        "Criterios de\nEvaluación",
        "Evidencia de\nAprendizaje",
        "Inst. de\nEvaluación"
    ]
    for col_idx, text in enumerate(headers):
        fill_cell(
            table_prop.cell(0, col_idx), 
            text, 
            bold=True, 
            bg_hex=BG_HEADER_HEX, 
            text_color_rgb=RGBColor(255, 255, 255), 
            size_pt=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        
    # Datos de Fila 1
    row1_data = [
        "[Competencia del Área]\n\n[Capacidades asociadas]",
        "[Desempeño precisado según el grado]",
        "[Qué se evaluará específicamente]",
        "[El producto o actuación del estudiante]",
        "[Lista de cotejo / Rúbrica]"
    ]
    for col_idx, text in enumerate(row1_data):
        fill_cell(table_prop.cell(1, col_idx), text, italic=True, size_pt=9.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4) # Separador
    
    # Enfoques transversales
    add_key_val_p(doc, "Enfoques Transversales Priorizados: ", "[Ej. Enfoque Ambiental / Enfoque Orientación al bien común]")
    add_key_val_p(doc, "Acciones observables: ", "[Descripción de la actitud de los estudiantes]", space_after=12)

    # IV. PREPARACIÓN DE LA SESIÓN DE APRENDIZAJE
    add_section_heading(doc, "IV. PREPARACIÓN DE LA SESIÓN DE APRENDIZAJE")
    
    prep_table = doc.add_table(rows=1, cols=2)
    set_table_borders(prep_table)
    set_col_widths(prep_table, [Inches(3.25), Inches(3.25)])
    
    fill_cell(
        prep_table.cell(0, 0), 
        "¿Qué necesitamos hacer antes de la sesión?\n\n[Insertar requerimientos previos]", 
        bold=False, 
        size_pt=10
    )
    fill_cell(
        prep_table.cell(0, 1), 
        "¿Qué recursos o materiales se utilizarán?\n\n[Ej. Fichas, proyector, pizarra, material concreto]", 
        bold=False, 
        size_pt=10
    )
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # V. MOMENTOS DE LA SESIÓN
    add_section_heading(doc, "V. MOMENTOS DE LA SESIÓN")
    
    # Inicio
    p_inicio = doc.add_paragraph()
    p_inicio.paragraph_format.space_before = Pt(8)
    p_inicio.paragraph_format.space_after = Pt(4)
    r = p_inicio.add_run("INICIO (Tiempo estimado: [15 min])")
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(46, 117, 89) # Verde oscuro
    
    add_key_val_p(doc, "• Motivación y Saberes Previos: ", "[Escribir cómo se captará el interés y qué preguntas se harán sobre experiencias pasadas].", bullet=False, space_after=3)
    add_key_val_p(doc, "• Conflicto Cognitivo: ", "[Planteamiento del reto o problema inicial].", bullet=False, space_after=3)
    add_key_val_p(doc, "• Propósito y Organización: ", "[Declaración explícita de lo que aprenderán hoy y cómo serán evaluados].", bullet=False, space_after=8)

    # Desarrollo
    p_des = doc.add_paragraph()
    p_des.paragraph_format.space_before = Pt(8)
    p_des.paragraph_format.space_after = Pt(4)
    r = p_des.add_run("DESARROLLO (Tiempo estimado: [60 min])")
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(74, 144, 226) # Azul
    
    p_nota = doc.add_paragraph()
    p_nota.paragraph_format.space_after = Pt(4)
    r_nota = p_nota.add_run("(Nota: Configurar los procesos didácticos según la naturaleza pedagógica del área curricular)")
    r_nota.italic = True
    r_nota.font.name = 'Calibri'
    r_nota.font.size = Pt(9.5)
    
    add_key_val_p(doc, "• Proceso Didáctico 1: ", "[Gestión y acompañamiento del aprendizaje / Planteamiento del problema o situación]", bullet=False, space_after=3)
    add_key_val_p(doc, "• Proceso Didáctico 2: ", "[Búsqueda y ejecución de estrategias por parte de los alumnos]", bullet=False, space_after=3)
    add_key_val_p(doc, "• Proceso Didáctico 3: ", "[Socialización de representaciones, formalización y reflexión del nuevo saber]", bullet=False, space_after=8)

    # Cierre
    p_cierre = doc.add_paragraph()
    p_cierre.paragraph_format.space_before = Pt(8)
    p_cierre.paragraph_format.space_after = Pt(4)
    r = p_cierre.add_run("CIERRE (Tiempo estimado: [15 min])")
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(192, 0, 0) # Rojo
    
    add_key_val_p(doc, "• Metacognición: ", "[Preguntas clave para el estudiante: ¿Qué aprendimos hoy? ¿Cómo lo aprendimos? ¿Para qué nos servirá?].", bullet=False, space_after=3)
    add_key_val_p(doc, "• Evaluación Formativa: ", "[Breve transferencia o sistematización final].", bullet=False, space_after=12)

    # Guardar
    doc.save(output_path)
    print(f"Plantilla de sesión creada en: {output_path}")


def generate_unidad_template(output_path):
    """Genera la Plantilla 2: Unidad de Aprendizaje."""
    doc = Document()
    
    # Configurar márgenes de página (1 pulgada = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Título Principal
    add_main_title(doc, "UNIDAD DE APRENDIZAJE")

    # I. DATOS INFORMATIVOS
    add_section_heading(doc, "I. DATOS INFORMATIVOS")
    
    # Tabla elegante para Datos Informativos (4 filas x 2 columnas)
    info_table = doc.add_table(rows=4, cols=2)
    set_table_borders(info_table)
    set_col_widths(info_table, [Inches(3.25), Inches(3.25)])
    
    # Llenar datos informativos
    fill_cell(info_table.cell(0, 0), "DRE / UGEL: [Especificar DRE y UGEL correspondiente]", bold=True)
    fill_cell(info_table.cell(0, 1), "Institución Educativa: [Nombre de la escuela]", bold=True)
    fill_cell(info_table.cell(1, 0), "Ciclo / Grado / Sección: [Ej. Ciclo IV - 4° de Primaria]", bold=True)
    fill_cell(info_table.cell(1, 1), "Director / Subdirector: [Nombres de las autoridades]", bold=True)
    fill_cell(info_table.cell(2, 0), "Docente Responsable: [Nombre del maestro]", bold=True)
    fill_cell(info_table.cell(2, 1), "Periodo de Ejecución: [Ej. Bimestre I / Trimestre I]", bold=True)
    fill_cell(info_table.cell(3, 0), "Duración Estimada: [Ej. Del 09 de marzo al 03 de abril]", bold=True)
    fill_cell(info_table.cell(3, 1), "", bold=True) # Celda vacía

    # II. TÍTULO DE LA UNIDAD DE APRENDIZAJE
    add_section_heading(doc, "II. TÍTULO DE LA UNIDAD DE APRENDIZAJE")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('[Escribir frase retadora y contextualizada, ej: "Asumimos desafíos para convivir armónicamente en un aula segura"]')
    run.font.italic = True
    run.font.name = 'Calibri'
    run.font.size = Pt(11)

    # III. SITUACIÓN SIGNIFICATIVA
    add_section_heading(doc, "III. SITUACIÓN SIGNIFICATIVA")
    add_key_val_p(doc, "Contexto: ", "[Descripción de la realidad local, problemática, necesidad o intereses de los estudiantes en su entorno].", space_after=4)
    add_key_val_p(doc, "Reto o Desafío: ", "[Preguntas movilizadoras y complejas, ej: ¿Cómo podemos organizar nuestra aula para trabajar en equipo? ¿Qué normas nos ayudarán?].", space_after=4)
    add_key_val_p(doc, "Producto Principal: ", "[El resultado final tangible o intangible, ej: Cartel de acuerdos de convivencia / Álbum comunal].", space_after=12)

    # IV. MATRIZ DE PROPÓSITOS DE APRENDIZAJE
    add_section_heading(doc, "IV. MATRIZ DE PROPÓSITOS DE APRENDIZAJE (Propósitos, Criterios y Actividades)")
    
    # Tabla de 5 columnas para la matriz
    # Total ancho = 6.5 pulgadas (1.0, 1.4, 1.4, 1.4, 1.3)
    table_matriz = doc.add_table(rows=3, cols=5)
    set_table_borders(table_matriz)
    set_col_widths(table_matriz, [Inches(1.0), Inches(1.4), Inches(1.4), Inches(1.4), Inches(1.3)])
    
    # Cabecera
    headers = [
        "Área",
        "Competencia /\nCapacidad",
        "Desempeños\ndel Grado",
        "Criterios de\nEvaluación",
        "Producción /\nActuación"
    ]
    for col_idx, text in enumerate(headers):
        fill_cell(
            table_matriz.cell(0, col_idx), 
            text, 
            bold=True, 
            bg_hex=BG_HEADER_HEX, 
            text_color_rgb=RGBColor(255, 255, 255), 
            size_pt=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER
        )
        
    # Datos de Fila 1
    row1_data = [
        "[Área]",
        "[Insertar Competencia]",
        "[Texto oficial del CNEB]",
        "[Indicador de logro esperado]",
        "[Evidencia medible]"
    ]
    for col_idx, text in enumerate(row1_data):
        fill_cell(table_matriz.cell(1, col_idx), text, italic=True, size_pt=9.5)
        
    # Datos de Fila 2 (Copia para mostrar multiplicidad)
    for col_idx, text in enumerate(row1_data):
        fill_cell(table_matriz.cell(2, col_idx), text, italic=True, size_pt=9.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Enfoques transversales
    add_key_val_p(doc, "Enfoques Transversales: ", "[Enfoque de Derechos / Enfoque Intercultural / etc.]", space_after=12)

    # V. SECUENCIA CONFIGURATIVA DE SESIONES
    add_section_heading(doc, "V. SECUENCIA CONFIGURATIVA DE SESIONES (Cronograma Modular)")
    
    add_key_val_p(doc, "• Sesión 1 ([Duración]): ", "Título: [Nombre de la Sesión 1] | Breve descripción del propósito pedagógico.", bullet=False, space_after=3)
    add_key_val_p(doc, "• Sesión 2 ([Duración]): ", "Título: [Nombre de la Sesión 2] | Breve descripción del propósito pedagógico.", bullet=False, space_after=3)
    add_key_val_p(doc, "• Sesión 3 ([Duración]): ", "Título: [Nombre de la Sesión 3] | Breve descripción del propósito pedagógico.", bullet=False, space_after=12)

    # VI. MATERIALES BÁSICOS Y RECURSOS A UTILIZAR
    add_section_heading(doc, "VI. MATERIALES BÁSICOS Y RECURSOS A UTILIZAR")
    p_mat1 = doc.add_paragraph(style='List Bullet')
    p_mat1.paragraph_format.space_after = Pt(3)
    r1 = p_mat1.add_run("[Textos escolares y cuadernos de trabajo distribuidos por el MINEDU]")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(11)
    
    p_mat2 = doc.add_paragraph(style='List Bullet')
    p_mat2.paragraph_format.space_after = Pt(12)
    r2 = p_mat2.add_run("[Plataformas digitales, herramientas de IA y recursos multimedia]")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(11)

    # VII. REFLEXIONES SOBRE LOS APRENDIZAJES
    add_section_heading(doc, "VII. REFLEXIONES SOBRE LOS APRENDIZAJES (Para el docente)")
    add_key_val_p(doc, "• ¿Qué avances tuvieron mis estudiantes? ", "[Espacio de balance]", space_after=3)
    add_key_val_p(doc, "• ¿Qué dificultades tuvieron mis estudiantes? ", "[Espacio de balance]", space_after=12)

    # Guardar
    doc.save(output_path)
    print(f"Plantilla de unidad creada en: {output_path}")


if __name__ == "__main__":
    # Rutas relativas para colocar en frontend/public
    public_dir = os.path.join("..", "frontend", "public")
    if not os.path.exists(public_dir):
        # Intentar ruta local de desarrollo
        public_dir = "frontend/public"
        
    if not os.path.exists(public_dir):
        # Fallback absoluta si es necesario
        public_dir = r"c:\Users\PC\Desktop\Avendia\frontend\public"
        
    os.makedirs(public_dir, exist_ok=True)
    
    clase_path = os.path.join(public_dir, "plantilla_clase_minedu.docx")
    unidad_path = os.path.join(public_dir, "plantilla_unidad_aprendizaje.docx")
    
    generate_clase_template(clase_path)
    generate_unidad_template(unidad_path)
